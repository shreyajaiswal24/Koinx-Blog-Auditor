"""Generate color-coded Excel audit reports using openpyxl."""

import logging
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config.settings import OUTPUT_DIR

logger = logging.getLogger(__name__)

# Color scheme
PRIORITY_FILLS = {
    "high": PatternFill(start_color="FF4444", end_color="FF4444", fill_type="solid"),
    "medium": PatternFill(start_color="FFB844", end_color="FFB844", fill_type="solid"),
    "low": PatternFill(start_color="44BB44", end_color="44BB44", fill_type="solid"),
}

STATUS_FILLS = {
    "new": PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid"),
    "previously_identified": PatternFill(start_color="87CEEB", end_color="87CEEB", fill_type="solid"),
    "resolved": PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid"),
}

HEADER_FILL = PatternFill(start_color="2B579A", end_color="2B579A", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
LINK_FONT = Font(color="0563C1", underline="single")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


class ExcelGenerator:
    """Generates Excel audit reports with multiple sheets."""

    def generate(
        self,
        findings: List[Dict[str, Any]],
        run_stats: Dict[str, Any],
        tax_references: List[Dict[str, str]],
    ) -> Path:
        """Generate the full Excel report.

        Args:
            findings: List of finding dicts with all metadata
            run_stats: Dict with run_id, started_at, total_posts, api_stats
            tax_references: List of tax reference dicts

        Returns:
            Path to the generated Excel file
        """
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filepath = OUTPUT_DIR / f"tax_audit_report_{timestamp}.xlsx"

        wb = Workbook()

        # Sheet 1: Summary
        self._create_summary_sheet(wb, run_stats, findings)

        # Sheet 2: Findings (main sheet)
        self._create_findings_sheet(wb, findings)

        # Sheet 3: Sources
        self._create_sources_sheet(wb, tax_references)

        # Remove default empty sheet if extra
        if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
            del wb["Sheet"]

        wb.save(str(filepath))
        logger.info(f"Report saved to {filepath}")
        return filepath

    def _create_summary_sheet(
        self, wb: Workbook, run_stats: Dict[str, Any], findings: List[Dict[str, Any]]
    ):
        """Create the Summary overview sheet."""
        ws = wb.create_sheet("Summary", 0)

        # Title
        ws.merge_cells("A1:D1")
        ws["A1"] = "KoinX Blog Tax Content Audit Report"
        ws["A1"].font = Font(bold=True, size=16)

        # Run info
        rows = [
            ("Run ID", run_stats.get("run_id", "N/A")),
            ("Timestamp", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("Total Posts Audited", run_stats.get("total_posts", 0)),
            ("Total Findings", len(findings)),
            ("High Priority", sum(1 for f in findings if f.get("priority") == "high")),
            ("Medium Priority", sum(1 for f in findings if f.get("priority") == "medium")),
            ("Low Priority", sum(1 for f in findings if f.get("priority") == "low")),
            ("New Findings", sum(1 for f in findings if f.get("status") == "new")),
            ("Previously Identified", sum(1 for f in findings if f.get("status") == "previously_identified")),
            ("Resolved", sum(1 for f in findings if f.get("status") == "resolved")),
            ("API Calls", run_stats.get("api_stats", {}).get("total_requests", 0)),
            ("Total Tokens Used", run_stats.get("api_stats", {}).get("total_tokens", 0)),
        ]

        for i, (label, value) in enumerate(rows, start=3):
            ws[f"A{i}"] = label
            ws[f"A{i}"].font = Font(bold=True)
            ws[f"B{i}"] = value

        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 20

    def _create_findings_sheet(self, wb: Workbook, findings: List[Dict[str, Any]]):
        """Create the main Findings sheet with color coding."""
        ws = wb.create_sheet("Findings")

        headers = [
            "Blog Link",
            "Blog Title",
            "Section",
            "Current Text (Exact Quote)",
            "Issue Type",
            "Description",
            "Suggested Updated Text",
            "Source",
            "Confidence",
            "Priority",
            "Status",
        ]

        # Write headers
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = THIN_BORDER

        # Freeze header row
        ws.freeze_panes = "A2"

        # Sort findings: high first, then medium, then low
        priority_order = {"high": 0, "medium": 1, "low": 2}
        sorted_findings = sorted(
            findings, key=lambda f: priority_order.get(f.get("priority", "low"), 3)
        )

        # Write findings
        for row_idx, finding in enumerate(sorted_findings, start=2):
            # Blog Link (hyperlinked)
            url = finding.get("blog_url", "")
            cell = ws.cell(row=row_idx, column=1)
            if url:
                cell.value = url
                cell.hyperlink = url
                cell.font = LINK_FONT
            else:
                cell.value = "N/A"

            ws.cell(row=row_idx, column=2, value=finding.get("blog_title", ""))
            ws.cell(row=row_idx, column=3, value=finding.get("section_heading", ""))
            ws.cell(row=row_idx, column=4, value=finding.get("exact_quote", ""))
            ws.cell(row=row_idx, column=5, value=finding.get("issue_type", ""))
            ws.cell(row=row_idx, column=6, value=finding.get("description", ""))
            ws.cell(row=row_idx, column=7, value=finding.get("suggested_update", ""))
            ws.cell(row=row_idx, column=8, value=finding.get("source", ""))
            ws.cell(row=row_idx, column=9, value=finding.get("confidence", 0))

            # Priority cell with color
            priority = finding.get("priority", "low")
            priority_cell = ws.cell(row=row_idx, column=10, value=priority.upper())
            if priority in PRIORITY_FILLS:
                priority_cell.fill = PRIORITY_FILLS[priority]
                priority_cell.font = Font(bold=True, color="FFFFFF")

            # Status cell with color
            status = finding.get("status", "new")
            status_cell = ws.cell(row=row_idx, column=11, value=status)
            if status in STATUS_FILLS:
                status_cell.fill = STATUS_FILLS[status]

            # Apply borders and wrap text
            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col)
                cell.border = THIN_BORDER
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        # Set column widths
        col_widths = [40, 35, 20, 50, 18, 50, 50, 35, 12, 10, 18]
        for i, width in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width

        # Add auto-filter
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(sorted_findings) + 1}"

    def _create_sources_sheet(self, wb: Workbook, tax_references: List[Dict[str, str]]):
        """Create the Sources reference sheet."""
        ws = wb.create_sheet("Tax Law Sources")

        headers = ["Topic", "Content", "Source", "URL"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.border = THIN_BORDER

        ws.freeze_panes = "A2"

        for row_idx, ref in enumerate(tax_references, start=2):
            ws.cell(row=row_idx, column=1, value=ref.get("topic", ""))
            ws.cell(row=row_idx, column=2, value=ref.get("content", "")[:2000])
            ws.cell(row=row_idx, column=3, value=ref.get("source", ""))

            url = ref.get("url", "")
            cell = ws.cell(row=row_idx, column=4)
            if url:
                cell.value = url
                cell.hyperlink = url
                cell.font = LINK_FONT
            else:
                cell.value = ""

            for col in range(1, 5):
                ws.cell(row=row_idx, column=col).border = THIN_BORDER
                ws.cell(row=row_idx, column=col).alignment = Alignment(
                    wrap_text=True, vertical="top"
                )

        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 80
        ws.column_dimensions["C"].width = 40
        ws.column_dimensions["D"].width = 50
